"""Generate a plain, ATS-safe .docx resume: single column, no tables,
no images, no text boxes, standard section headings, Calibri 11pt.
Run with: python scripts/generate_ats_docx.py
"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "ats")
OUT_PATH = os.path.join(OUT_DIR, "Muhammad_Zaim_Rusli_Resume.docx")

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def heading(text):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)

def para(text, bold=False, size=11, align_center=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def bullet(text):
    doc.add_paragraph(text, style="List Bullet")

def entry_header(role, org, dates):
    p = doc.add_paragraph()
    r1 = p.add_run(f"{role} — {org}")
    r1.bold = True
    p.add_run(f"\t{dates}")
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Pt(430))

para("Muhammad Zaim Bin Rusli", bold=True, size=16, align_center=True)
para("Senior System Validation & Silicon Debug Engineer", size=12, align_center=True)
para("mzaimrusli@gmail.com | +60 16-610 4339 | linkedin.com/in/muhammad-zaim-rusli-805715108 | Penang, Malaysia", align_center=True)

heading("Professional Summary")
para(
    "Senior System Validation and Silicon Debug Engineer with nearly 10 years of experience spanning "
    "semiconductor manufacturing, ATE test development, pre-silicon verification, post-silicon validation, "
    "System Level Test (SLT) and data center platform qualification across Intel Xeon and AMD EPYC server processors."
)
para(
    "Proven expertise in system-level validation, production readiness qualification, silicon bring-up, "
    "failure analysis, yield improvement, debug, hardware-software-firmware integration and manufacturing "
    "test optimization. Experienced in developing automated validation scripts using Python and Shell "
    "scripting while driving cross-domain root cause analysis across silicon, firmware, BIOS, operating "
    "systems and platform hardware."
)
para(
    "Unique background spanning wafer sort, package test, ATE, SLT, pre-silicon verification and post-silicon "
    "validation enables strong correlation across manufacturing, lab validation and field deployment environments."
)

heading("System Level Test and Manufacturing Highlights")
for b in [
    "Nearly 10 years of experience across semiconductor manufacturing test, ATE, post-silicon validation and production qualification.",
    "Hands-on experience supporting High Volume Manufacturing (HVM) environments.",
    "Developed and maintained automated test scripts for validation and production readiness.",
    "Expertise correlating manufacturing failures, system-level failures and silicon root causes.",
    "Extensive experience debugging issues across hardware, firmware, BIOS and operating system boundaries.",
    "Strong background in server platform bring-up including CPU, memory, PCIe devices, NICs and OpenBMC infrastructures.",
    "Experience supporting release and manufacturing decisions based on validation readiness.",
    "Proven ability to identify validation and test coverage gaps through structured failure analysis.",
    "Experienced in SLT, Manufacturing Readiness, Production Qualification, Yield Improvement, "
    "Failure Correlation, Tester Correlation, Root Cause Analysis, Platform Bring-up.",
]:
    bullet(b)

heading("Technical Skills")
para("Architecture & Platforms: AMD EPYC, Intel Xeon, x86 Architecture, Server Platforms, Datacenter Infrastructure, PCIe Gen5/Gen6, DDR4/DDR5, BIOS/UEFI, OpenBMC")
para("Programming & Automation: Python, PyTest, Shell/Bash, C/C++, SystemVerilog, UVM, Simics Automation")
para("Debug & Diagnostics: Machine Check Exceptions (MCE), WHEA, BERT, Error Injection, Status Registers, Scan Dumps, dmesg, serial/UART console, SPI flash reprogramming, JTAG")
para("Tools: PRISM, HEXA, Git, JIRA, Confluence, Simics, Verdi, DediProg, Linux, Windows, Xilinx Vivado")

heading("Professional Experience")

entry_header("Functional Validation Engineer", "Advanced Micro Devices (AMD), Penang, Malaysia", "June 2024 - Present")
for b in [
    "Bring up new AMD EPYC server platforms end-to-end from receipt of engineering systems, including HW installation (DIMM, CPU), troubleshooting (NIC, LAN, component), firmware flashing (BIOS, BMC, FPGA), OS provisioning and OpenBMC configuration, to establish a stable boot baseline ahead of RAS test execution.",
    "Triage platform bring-up failures on new silicon, isolating root cause across DIMM faults, CPLD corruption, outdated firmware recipes (BMC/BIOS/FPGA version mismatches), physically damaged components, leakage and device enumeration failures preventing system boot.",
    "Validate RAS error detection, classification and recovery behavior for XGMI flit errors, PCIe AER (Correctable/Non-Fatal/Fatal) and DRAM ECC (single-bit/multi-bit) on server platforms.",
    "Analyze error logs and traces across dmesg, BERT, register dump, Sherlock, WinDbg and DebugView to confirm errors are correctly detected, logged and surfaced to the OS/firmware without being silently dropped or misclassified.",
    "Develop and maintain a pytest-based PCIe AER RAS validation suite covering correctable and uncorrectable (fatal/non-fatal) error types such as LCRC, ECRC and poisoned TLP, across OS-First and Firmware-First (SMI/MPRAS) reporting modes, with Downstream Port Containment on/off, using both ACPI EINJ and direct register-level (SMN) error injection.",
    "Implement register-level PCI configuration space overrides (via setpci) to force root-port DPC and error-mask register state to match intended test conditions when BIOS defaults diverged, restoring original state after each test.",
    "Leverage AI coding assistants (Claude) to accelerate development and maintenance of Python-based RAS validation scripts, freeing up engineering time for validation, debug and silicon root-cause analysis.",
    "Own regression across platform configurations against evolving FW builds within the PRISM framework, triaging failures to root cause across script, framework, platform or silicon.",
    "Collaborate cross-functionally with architecture, firmware, BIOS and HW teams to close out complex system-level RAS issues and improve validation coverage, tracking issues in Jira and maintaining debug/validation documentation in Confluence.",
]:
    bullet(b)

entry_header("System Validation Engineer", "Intel Microelectronics, Kulim, Malaysia", "August 2019 - May 2024")
para("Post-Silicon System Validation", bold=True)
for b in [
    "Improved and executed structured test plans and silicon qualification activities on the Intel Xeon server platform, defining system-level validation methodologies that directly supported production release decisions.",
    "Developed and maintained Python and Shell-based automation within the HEXA validation framework to improve regression efficiency and validation scalability.",
    "Validated ACM authentication and SEAM Loader boot/initialization flows for TDX, confirming correct TDX module measurement and loading across BIOS and microcode updates.",
    "Validated SGX and TDX attestation flows, confirming enclave/TD quotes correctly reflect genuine platform and measurement state.",
    "Validated concurrency and stress testing involving hundreds of simultaneous enclaves/Trust Domains to characterize silicon reliability and validate isolation integrity under production-like load.",
    "Performed root cause analysis across hardware, BIOS, firmware and OS layers, collaborating with architecture teams to resolve manufacturing-blocking and customer-impacting silicon defects.",
]:
    bullet(b)
para("Pre-Silicon Validation and Emulation", bold=True)
for b in [
    "Responsible for SoC functional validation of Intel Xeon server processors in Simics Virtual Platform and FPGA emulation environments ahead of silicon availability.",
    "Developed JSON-configurable automation to drive boot sequences and verify TDX/SGX-related register state against specification, enabling fast scenario experimentation without code changes.",
    "Debugged cases where expected registers failed to reach their specified state, tracing root cause across initialization ordering and distinguishing genuine design issues from virtual-platform model limitations.",
]:
    bullet(b)
para("Client and Embedded Validation", bold=True)
for b in [
    "Performed functional validation on client/laptop platform SoCs using Simics Virtual Platform, developing automation (legacy in-house framework) to drive the full flow from system boot through BIOS configuration changes to device/register verification on the system under test.",
    "Collaborated with firmware and software teams to resolve integration issues and improve platform stability.",
]:
    bullet(b)

entry_header("IP Verification Engineer", "Intel Microelectronics, Penang, Malaysia", "March 2018 - August 2019")
for b in [
    "Verified a soft IP block's functionality using SystemVerilog/UVM, improving and extending constrained-random test scenarios that uncovered handshake failures and timing-window violations in signal insertion, broadening test coverage to catch them.",
    "Performed waveform-level root cause analysis using Verdi to debug the protocol and timing issues identified during verification.",
    "Collaborated with RTL design and architecture teams to resolve protocol and integration issues or limitations ahead of release.",
]:
    bullet(b)

entry_header("Test and Product Engineer", "Malaysia Microelectronic Solutions (MyMS), Selangor, Malaysia", "August 2017 - March 2018")
for b in [
    "Developed and debugged ATE test programs on Credence Kalos testers for production IC test, including a smart card (payphone card) product for an overseas telecom customer.",
    "Performed wafer probing on an Electroglas EG4090µ wafer prober loaned from MIMOS, diagnosing probe-related hardware issues including prober die-indexing/mapping mismatches and resolving them in coordination with MIMOS, while supporting yield analysis and electrical characterization for production release readiness.",
]:
    bullet(b)

entry_header("Graduate Engineer", "NXP Semiconductors, Selangor, Malaysia", "September 2016 - July 2017")
for b in [
    "Enhanced Teradyne J750 production test programs for automotive MCU devices as part of the NVM (non-volatile memory) test team.",
    "Performed device characterization and tester-to-golden-unit correlation to validate test accuracy, supporting failure analysis, yield improvement and production qualification.",
]:
    bullet(b)

heading("Education")
para("Bachelor of Engineering (Hons.), Electronic Engineering — Universiti Teknologi MARA (UiTM), 2016")

os.makedirs(OUT_DIR, exist_ok=True)
doc.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")
